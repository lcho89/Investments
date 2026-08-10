#!/usr/bin/env python3
"""Print the tab names in the downloaded workbook, so sources.json can be filled in."""
import sys
from pathlib import Path
from openpyxl import load_workbook
p = Path(__file__).resolve().parent.parent / "portfolio" / "exports" / "workbook.xlsx"
if not p.exists():
    sys.exit(f"No workbook at {p}. Run: python3 scripts/refresh_holdings.py (it will fetch it)")
for name in load_workbook(p, read_only=True).sheetnames:
    print(name)
